import io
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import streamlit as st

st.set_page_config(page_title="木谷さんW8LY専用エキスパート整理ツール", page_icon="📋", layout="wide")

# 指定のカラーコード（背景: #355E3B, ボックス: #F5EFD6, 文字: #895129）を適用したCSS
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Zen+Kaku+Gothic+New:wght@700;900&display=swap');

    .stApp {
        background-color: #355E3B !important;
        background-image: none !important;
    }
    
    .main-card {
        background-color: #F5EFD6 !important;
        border: 2px solid #895129 !important;
        border-radius: 24px;
        padding: 40px 35px;
        box-shadow: 0 15px 35px rgba(0, 0, 0, 0.25);
        color: #895129 !important;
        text-align: center;
        max-width: 700px;
        margin: 10px auto 25px auto;
    }

    .badge-theme {
        background-color: #895129 !important;
        color: #F5EFD6 !important;
        padding: 6px 20px;
        border-radius: 30px;
        font-size: 11px;
        font-weight: 800;
        letter-spacing: 2.5px;
        text-transform: uppercase;
        display: inline-block;
        margin-bottom: 14px;
    }

    h1.main-title {
        font-family: 'Zen Kaku Gothic New', 'Hiragino Sans', sans-serif;
        color: #895129 !important;
        font-size: 28px;
        font-weight: 900;
        margin: 0 0 12px 0;
        letter-spacing: -0.5px;
        -webkit-text-fill-color: #895129 !important;
    }

    p.sub-desc {
        color: #895129 !important;
        font-size: 13px;
        font-weight: 700;
        line-height: 1.6;
        margin-bottom: 0;
    }

    .stButton > button {
        background-color: #895129 !important;
        color: #F5EFD6 !important;
        border: none !important;
        padding: 14px 32px !important;
        border-radius: 50px !important;
        font-family: 'Zen Kaku Gothic New', sans-serif !important;
        font-size: 16px !important;
        font-weight: 900 !important;
        letter-spacing: 1.5px !important;
        box-shadow: 0 8px 20px rgba(0, 0, 0, 0.2) !important;
        transition: all 0.3s ease !important;
        width: 100% !important;
    }

    .stButton > button:hover {
        transform: translateY(-2px) !important;
        background-color: #704020 !important;
    }

    div[data-testid="stTextArea"] textarea {
        background-color: #F5EFD6 !important;
        color: #222222 !important;
        border: 2px solid #895129 !important;
        border-radius: 16px !important;
        font-size: 13.5px !important;
    }
    
    div[data-testid="stTextArea"] label {
        color: #F5EFD6 !important;
        font-weight: 800 !important;
    }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="main-card">
    <span class="badge-theme">W8LY SPECIAL TOOL</span>
    <h1 class="main-title">木谷さんW8LY専用エキスパート整理ツール</h1>
    <p class="sub-desc">メール本文をそのままコピペして貼り付けるだけで、メール送信用テキスト整形とエクセルデータの作成を同時に行います。</p>
</div>
""", unsafe_allow_html=True)

input_text = st.text_area("▼ 送信されてきたメール文面をここに貼り付けてください", height=280, placeholder="Hi Yui, ... から始まるメールテキストをそのままペースト")

def parse_email_text(raw_text):
    pattern = r'(?=(?:^|\n)#\d+(?:\.\d+)?\s*-)'
    chunks = re.split(pattern, raw_text)
    
    parsed_experts = []
    
    for chunk in chunks:
        chunk_str = chunk.strip()
        if not chunk_str or not re.search(r'#\d+(?:\.\d+)?', chunk_str):
            continue
            
        header_match = re.search(r'#(\d+(?:\.\d+)?)\s*-\s*([^-]+?)\s*-\s*(.+?)(?=\n|$)', chunk_str)
        if not header_match:
            continue
            
        number = header_match.group(1).strip()
        name = header_match.group(2).strip()
        title = header_match.group(3).strip()
        
        lines = chunk_str.split('\n')
        summary_lines = []
        screened_lines = []
        emp_lines = []
        avail_lines = []
        
        in_screened = False
        in_emp = False
        in_avail = False
        
        for line in lines[1:]:
            l = line.strip()
            if not l:
                continue
                
            if re.match(r'\[\s*(?:Screened|Re-screened).*?\]', l, re.IGNORECASE) or re.match(r'Screened\s+\d', l, re.IGNORECASE):
                in_screened = True
                in_emp = False
                in_avail = False
                screened_lines.append(l)
                continue
            elif "Employment History:" in l:
                in_screened = False
                in_emp = True
                in_avail = False
                continue
            elif "Availability:" in l:
                in_screened = False
                in_emp = False
                in_avail = True
                continue
            elif any(k in l for k in ["This specialist has not yet provided any availability", "Request Availability", "Book Now", "This specialist is based in", "Hourly Fee:"]):
                if in_avail and "Time Zone:" not in l:
                    pass
                continue
            elif any(k in l for k in ["Disclaimer:", "Important:", "Compliance Reminder", "Third Bridge"]):
                break
                
            if in_screened:
                screened_lines.append(l)
            elif in_emp:
                emp_lines.append(l)
            elif in_avail:
                if "Time Zone:" in l:
                    continue
                avail_lines.append(l)
            else:
                summary_lines.append(l)
                
        summary_text = "\n".join(summary_lines)
        qa_text = "\n".join(screened_lines)
        emp_text = "\n".join(emp_lines)
        
        if avail_lines:
            avail_text = "\n".join(avail_lines)
        else:
            avail_text = "回収中"
            
        parsed_experts.append({
            "number": number,
            "name": name,
            "title": title,
            "summary": summary_text,
            "qa": qa_text,
            "emp": emp_text,
            "availability": avail_text
        })
        
    return parsed_experts

def generate_formatted_email_text(experts):
    output_blocks = []
    for exp in experts:
        block = f"#{exp['number']} - {exp['name']} - {exp['title']}\n"
        if exp['summary']:
            block += f"\n{exp['summary']}\n"
        if exp['qa']:
            block += f"\n{exp['qa']}\n"
        
        block += f"\nAvailability:\n{exp['availability']}\n"
        block += "-" * 50
        output_blocks.append(block)
        
    return "\n\n".join(output_blocks)

def generate_excel_bytes(experts):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "全員一覧"
    
    headers = ["Number", "Name", "Relevant Titles", "Relevant experience", "Employment History", "Availability"]
    
    font_bold = Font(name="Meiryo UI", size=9, bold=True)
    font_regular = Font(name="Meiryo UI", size=9)
    fill_header = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )
    
    ws["A1"] = "ThirdBridge"
    ws["A1"].font = font_bold
    
    for c_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c_idx, value=h_text)
        cell.font = font_bold
        cell.fill = fill_header
        cell.alignment = Alignment(vertical="top")
        
    for r_idx, exp in enumerate(experts, 4):
        exp_full_text = exp['summary']
        if exp['qa']:
            exp_full_text += f"\n\n{exp['qa']}"
            
        row_vals = [
            exp['number'],
            exp['name'],
            exp['title'],
            exp_full_text,
            exp['emp'],
            exp['availability']
        ]
        
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            
            if c_idx == 1:
                try:
                    cell.value = float(val) if "." in str(val) else int(val)
                    cell.number_format = '0.0'
                except ValueError:
                    pass
                cell.alignment = Alignment(horizontal="center", vertical="top")
                
    col_widths = [10, 18, 38, 60, 50, 30]
    for c_i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(c_i)].width = w
        
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

if input_text:
    parsed_experts = parse_email_text(input_text)
    
    if parsed_experts:
        st.success(f"✨ {len(parsed_experts)}名のエキスパート情報を正常に抽出しました！")
        
        st.markdown("### 📧 メールコピペ用テキスト")
        email_formatted_text = generate_formatted_email_text(parsed_experts)
        st.code(email_formatted_text, language="text")
        
        excel_data = generate_excel_bytes(parsed_experts)
        st.download_button(
            label="📥 エクセルファイルをダウンロード",
            data=excel_data,
            file_name="W8LY_Expert_List.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        st.warning("エキスパート情報（#番号 で始まる記述）が検出されませんでした。入力内容をご確認ください。")
